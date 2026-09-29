"""sahtTatva operational pipeline — one command runs everything.

    python src/run_daily.py

Steps: ingest -> bias_correct -> blend (weights) -> verify -> snapshots
       -> alerts -> finalize (weight_maps.png, summary.xlsx)

All outputs land in outputs/. Every number comes from the live Open-Meteo
APIs (forecast + archive + geocoding); nothing is synthetic.
"""
import json
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))

import alerts
import bias_correct
import blend
import ingest
import snapshots
import verify
from config import (APP_NAME, BASE_DIR, CITIES, DATA_DIR, LEAD_BANDS, MODELS,
                    MODEL_LABELS, OUT_DIR, TIMEZONE, VARIABLES, VAR_LABELS)


def build_weight_maps():
    """weight_maps.png — which model is trusted where/when (PS asks for this)."""
    with open(DATA_DIR / "weights.json") as f:
        gw = json.load(f)["weights"]["global"]
    bands = list(LEAD_BANDS)
    fig, axes = plt.subplots(1, 3, figsize=(15, 4.5))
    palette = ["#5c80bc", "#7bc96f", "#e07a3f", "#9b59b6", "#e74c3c"]
    colors = {m: palette[i % len(palette)] for i, m in enumerate(MODELS)}
    for ax, var in zip(axes, VARIABLES):
        x = np.arange(len(bands))
        width = 0.22
        for i, model in enumerate(MODELS):
            vals = [gw[var][b][model] for b in bands]
            ax.bar(x + (i - 1) * width, vals, width,
                   label=MODEL_LABELS[model], color=colors[model])
        ax.set_xticks(x)
        ax.set_xticklabels(bands)
        ax.set_ylim(0, 1)
        ax.set_ylabel("Weight")
        ax.set_title(VAR_LABELS[var], fontsize=11)
        ax.legend(fontsize=8)
    fig.suptitle(f"{APP_NAME} — adaptive blending weights "
                 "(inverse test-RMSE, sum to 1)", fontsize=13, fontweight="bold")
    fig.tight_layout()
    fig.savefig(OUT_DIR / "weight_maps.png", dpi=130)
    plt.close(fig)
    print("Saved outputs/weight_maps.png")


def _style_sheet(ws, df):
    header_font = Font(bold=True, color="FFFFFF")
    from openpyxl.styles import PatternFill
    fill = PatternFill("solid", fgColor="2F5496")
    for c, col in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=c, value=col)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for c, col in enumerate(df.columns, 1):
        width = max(12, min(38, int(df[col].astype(str).str.len().max()
                                   if len(df) else 12) + 2))
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_summary_xlsx():
    """summary.xlsx — Cover, Training Cities, Weights, Verification, Alerts."""
    wb = Workbook()
    now = datetime.now(ZoneInfo(TIMEZONE)).strftime("%Y-%m-%d %H:%M IST")

    ws = wb.active
    ws.title = "Cover"
    cover = [
        f"{APP_NAME} — Hybrid AI-NWP Multi-Model Forecast Blending System",
        "SIH 2026 Problem Statement SIH26081 | Ministry of Earth Sciences (MoES)",
        f"Generated: {now}",
        "",
        "What this is:",
        "An adaptive system blending GFS, ICON and GEM forecasts with ML bias",
        "correction (gradient boosting on real model-analysis vs ERA5 pairs),",
        "a do-no-harm gate, and inverse-skill weighting per lead-time band.",
        "",
        "How to run:",
        "1. python -m venv venv && venv/bin/pip install -r requirements.txt",
        "2. python src/run_daily.py   # refresh all data + outputs",
        "3. streamlit run dashboard/app.py   # launch the live UP search app",
    ]
    for i, text in enumerate(cover, 1):
        c = ws.cell(row=i, column=1, value=text)
        if i <= 3:
            c.font = Font(bold=True, size=13 if i == 1 else 11)
    ws.column_dimensions["A"].width = 90

    with open(DATA_DIR / "training_stats.json") as f:
        stats = json.load(f)
    cities_df = pd.DataFrame([
        {"City": c, "Latitude": lat, "Longitude": lon,
         "Train days": stats["per_city_train_days"][c],
         "Test days": stats["per_city_test_days"][c]}
        for c, (lat, lon) in CITIES.items()
    ])
    ws2 = wb.create_sheet("Training Cities")
    for r, row in enumerate([cities_df.columns.tolist()] + cities_df.values.tolist(), 1):
        for c, v in enumerate(row, 1):
            ws2.cell(row=r, column=c, value=v)
    _style_sheet(ws2, cities_df)
    ws2.cell(row=len(cities_df) + 3, column=1,
             value=f"Pooled training rows (train split): {stats['training_rows_pooled']:,}")
    ws2.cell(row=len(cities_df) + 4, column=1,
             value="Source: Open-Meteo forecast API (model analyses) + ERA5 archive. Zero synthetic data.")

    wdf = pd.read_csv(DATA_DIR / "weights.csv")
    wdf = wdf[wdf["scope"] == "global"].drop(columns=["scope"])
    ws3 = wb.create_sheet("Weights (global)")
    for r, row in enumerate([wdf.columns.tolist()] + wdf.values.tolist(), 1):
        for c, v in enumerate(row, 1):
            ws3.cell(row=r, column=c, value=v)
    _style_sheet(ws3, wdf)

    vdf = pd.read_csv(OUT_DIR / "verification_table.csv")
    upv = vdf[vdf["city"] == "UP-AGGREGATED"].drop(columns=["city"])
    ws4 = wb.create_sheet("Verification (UP)")
    for r, row in enumerate([upv.columns.tolist()] + upv.values.tolist(), 1):
        for c, v in enumerate(row, 1):
            ws4.cell(row=r, column=c, value=v)
    _style_sheet(ws4, upv)

    cityv = vdf[vdf["city"] != "UP-AGGREGATED"]
    ws5 = wb.create_sheet("Verification (per city)")
    for r, row in enumerate([cityv.columns.tolist()] + cityv.values.tolist(), 1):
        for c, v in enumerate(row, 1):
            ws5.cell(row=r, column=c, value=v)
    _style_sheet(ws5, cityv)

    adf = pd.read_csv(OUT_DIR / "alerts.csv")
    ws6 = wb.create_sheet("Alerts")
    if len(adf):
        for r, row in enumerate([adf.columns.tolist()] + adf.values.tolist(), 1):
            for c, v in enumerate(row, 1):
                ws6.cell(row=r, column=c, value=v)
        _style_sheet(ws6, adf)
    else:
        ws6.cell(row=1, column=1, value="No extreme-weather alerts in the current 7-day window.")

    path = OUT_DIR / "summary.xlsx"
    wb.save(path)
    print(f"Saved {path}")


def main():
    print(f"===== {APP_NAME} daily pipeline =====")
    ingest.main()
    bias_correct.main()
    blend.main()
    verify.main()
    snapshots.main()
    alerts.main()
    build_weight_maps()
    build_summary_xlsx()
    # also copy the key blended snapshot table for convenience
    snap = pd.read_csv(OUT_DIR / "snapshots_blended.csv", parse_dates=["time"])
    snap.to_csv(OUT_DIR / "blended_forecast.csv", index=False)
    print("Saved outputs/blended_forecast.csv (12-city 7-day blended snapshots)")
    print("===== pipeline complete =====")


if __name__ == "__main__":
    main()
